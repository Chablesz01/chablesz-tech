import csv
import importlib.util
import shutil
import subprocess
import unittest
from pathlib import Path
from PySide6.QtCore import QObject,Signal,QTimer,Qt
from PySide6.QtGui import QImage,QColor,QPixmap
from PySide6.QtWidgets import QLabel
from PySide6.QtTest import QTest
import app as module
import test_ui
from bible_core import normalize_rows,search_rows
from projection import SlideCanvas
from browser_bridge import BrowserSession
from agenda_core import save_package,load_package_full
from web_media import validated_url,media_script

class DummyBrowser(QObject):
    tapped=Signal();message=Signal(str)
    attach=BrowserSession.attach
    def __init__(self,parent):
        super().__init__(parent);self.widget=QLabel('Prepared web page');self.stage=None
        self.settings={'preview':True,'fill':True,'volume':.8};self.closed=False;self.playback_time=42.5
        self.applied=0
    def apply_settings(self):self.applied+=1
    def close(self):self.closed=True;self.widget.hide()
    def fill(self,value):self.settings['fill']=value
    def media(self,*args):pass
    def snapshot(self):return QPixmap()

class V11Tests(unittest.TestCase):
    setUpClass=classmethod(test_ui.UITests.setUpClass.__func__)
    setUp=test_ui.UITests.setUp
    tearDown=test_ui.UITests.tearDown
    search=test_ui.UITests.search
    open_live=test_ui.UITests.open_live

    def test_bundled_editions_and_reference(self):
        for version,count in [('KJV',31102),('ASV',31086),('WEB',31102),('YLT',31102)]:
            with (Path(module.__file__).parent/'bundled_bibles'/(version+'.csv')).open(encoding='utf-8') as f:
                rows=normalize_rows(list(csv.DictReader(f)))
            self.assertEqual(len(rows),count);self.assertEqual(len({r['book'] for r in rows}),66)
            self.assertEqual(len(search_rows(rows,'John 3:16')),1)
            self.w.version.setCurrentText(version);self.search('John 3:16')
            self.assertIn('('+version+')',self.w.selected_verse()['reference'])

    def test_independent_images_headers_and_reference_position(self):
        images={}
        for profile,color in [('bible','#004400'),('song','#440000')]:
            image=QImage(160,90,QImage.Format_RGB32);image.fill(QColor(color))
            path=Path(self.temp.name)/(profile+'.png');image.save(str(path));images[profile]=str(path)
            self.w.preview.settings[profile+'_background_image']=str(path)
        self.w.preview.settings.update(bible_header_enabled=True,bible_header_text='RCCG JESUS THE CHAMPION',bible_reference_position='Top')
        self.search('John 3:16');self.qt.processEvents();self.w.preview.grab()
        canvas=self.w.preview.content_layout.itemAt(0).widget()
        self.assertIsInstance(canvas,SlideCanvas)
        self.assertEqual(canvas.image.toImage().pixelColor(0,0).name(),'#004400')
        self.assertLessEqual(canvas.header_rect.bottom(),canvas.reference_rect.top())
        self.assertLessEqual(canvas.reference_rect.bottom(),canvas.body_rect.top())
        self.w.preview.display({'kind':'text','role':'song','text':'A worship song','title':'Song'})
        self.qt.processEvents();canvas=self.w.preview.content_layout.itemAt(0).widget()
        self.assertEqual(canvas.image.toImage().pixelColor(0,0).name(),'#440000')
        package=Path(self.temp.name)/'service.chablesz'
        styles={profile+'_background_image':path for profile,path in images.items()}
        save_package(package,[{'kind':'text','text':'Slide','title':'Slide'}],styles)
        _,loaded=load_package_full(package,Path(self.temp.name)/'restored')
        self.assertTrue(all(Path(loaded[p+'_background_image']).exists() for p in images))

    def test_web_handoff_retains_instance_time_and_fill(self):
        item={'kind':'url','url':'https://www.youtube.com/watch?v=example','title':'Prepared web video'}
        self.w.items=[item];self.w.refresh_list(preview=False)
        session=DummyBrowser(self.w.preview)
        self.w.preview.display(item,web_session=session)
        self.w.send_preview_live();self.qt.processEvents()
        self.assertIs(self.w.live.web_session,session)
        self.assertFalse(session.closed);self.assertEqual(session.playback_time,42.5)
        self.assertTrue(session.settings['fill']);self.assertFalse(session.settings['preview'])
        self.assertIs(session.widget.parent(),self.w.live.content)
        self.w.live.blackout(True);self.w.live.blackout(False)
        self.assertIs(self.w.live.web_session,session);self.assertFalse(session.closed)

    def test_full_preview_returns_without_recreating_browser(self):
        item={'kind':'url','url':'https://example.com','title':'Web'}
        session=DummyBrowser(self.w.preview);self.w.preview.display(item,web_session=session)
        self.w.preview_size.setCurrentIndex(1);self.qt.processEvents()
        self.assertIs(self.w.preview.parent(),self.w.preview_popup)
        self.assertFalse(session.settings['tap_enabled'])
        self.w.preview_size.setCurrentIndex(0);self.qt.processEvents()
        self.assertIs(self.w.preview.parent(),self.w.preview_frame)
        self.assertTrue(session.settings['tap_enabled']);self.assertFalse(session.closed)

    def test_single_tap_video_moves_to_live(self):
        if not shutil.which('ffmpeg'):self.skipTest('ffmpeg fixture generator is unavailable')
        path=Path(self.temp.name)/'tap.mp4'
        subprocess.run(['ffmpeg','-loglevel','error','-f','lavfi','-i','color=c=blue:s=320x180:d=3',
                        '-c:v','mpeg4','-y',str(path)],check=True)
        item={'kind':'video','path':str(path),'title':'Tap video'}
        self.w.items=[item];self.w.refresh_list()
        QTest.qWait(150);self.w.preview.player.setPosition(600);self.w.preview.player.pause();QTest.qWait(30)
        video=self.w.preview.content_layout.itemAt(0).widget()
        QTest.mouseClick(video,Qt.LeftButton);QTest.qWait(100)
        self.assertEqual(self.w.current_live_index,0)
        self.assertEqual(self.w.live.current,item)
        self.assertTrue(self.w.preview.audio.isMuted())

    def test_real_powerpoint_conversion_keeps_ui_ticking(self):
        if not (shutil.which('soffice') or shutil.which('libreoffice')) or not importlib.util.find_spec('pptx'):
            self.skipTest('LibreOffice and the pptx fixture generator are required')
        from pptx import Presentation
        from pptx.util import Inches
        deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
        for number in range(8):
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(11),Inches(4))
            shape.text_frame.text='PowerPoint import test • Slide '+str(number+1)
        path=Path(self.temp.name)/'eight_slides.pptx';deck.save(path)
        ticks=[];timer=QTimer(self.w);timer.timeout.connect(lambda:ticks.append(1));timer.start(15)
        self.w.import_paths([str(path)])
        for _ in range(400):
            QTest.qWait(25)
            if self.w.import_worker is None:break
        timer.stop()
        self.assertIsNone(self.w.import_worker,'Conversion must finish within the test deadline')
        self.assertEqual(len(self.w.items),8)
        self.assertGreater(len(ticks),5)
        self.assertEqual([i['slide_number'] for i in self.w.items],list(range(1,9)))
        self.assertTrue(all(Path(i['path']).is_file() for i in self.w.items))

    def test_url_validation_and_control_scripts(self):
        self.assertEqual(validated_url('youtube.com'),'https://youtube.com')
        for invalid in ['https://user:password@example.com','https://','http://']:
            with self.assertRaises(ValueError):validated_url(invalid)
        self.assertIn('position:fixed',media_script('fill'))
        self.assertIn('.remove()',media_script('page'))
        self.assertIn('chableszWasPlaying',media_script('black'))
        self.assertIn('Repeat once',media_script('repeat','Repeat once'))
        self.assertIn('currentTime',media_script('seek',10))


    def test_enter_and_physical_keys_flow_through_passage(self):
        self.w.query.setFocus();self.w.query.setText('John 2:3-8')
        QTest.keyClick(self.w.query,Qt.Key_Return);self.qt.processEvents()
        self.assertEqual(len(self.w.bible_flow_indices()),6)
        self.assertIn('John 2:3',self.w.live.current['reference'])
        for key,verse in [(Qt.Key_PageDown,4),(Qt.Key_Right,5),(Qt.Key_Down,6),
                          (Qt.Key_PageUp,5),(Qt.Key_Left,4),(Qt.Key_Up,3)]:
            QTest.keyClick(self.w.results,key);self.qt.processEvents()
            self.assertIn('John 2:'+str(verse),self.w.live.current['reference'])
        self.assertEqual(len(self.w.items),6)

    def test_passage_boundaries_and_invalid_search_preserve_live(self):
        self.search('John 2:3-4');self.w.verse_live()
        self.w.items.append({'kind':'text','text':'Unrelated slide','title':'Other'})
        self.w.previous_live();self.assertIn('John 2:3',self.w.live.current['reference'])
        self.w.next_live();self.w.next_live();self.assertIn('John 2:4',self.w.live.current['reference'])
        self.w.query.setText('John 999:1');QTest.keyClick(self.w.query,Qt.Key_Return)
        self.assertIn('John 2:4',self.w.live.current['reference'])
        self.assertEqual(len(self.w.items),3)

    def test_typing_keys_and_live_window_navigation(self):
        self.search('John 2:3-5');self.w.verse_live()
        self.w.query.setFocus();self.w.query.setCursorPosition(4)
        QTest.keyClick(self.w.query,Qt.Key_Left)
        self.assertEqual(self.w.query.cursorPosition(),3)
        self.assertIn('John 2:3',self.w.live.current['reference'])
        QTest.keyClick(self.w.live,Qt.Key_PageDown)
        self.assertIn('John 2:4',self.w.live.current['reference'])

    def test_new_passage_and_selected_verse_start(self):
        self.search('John 2:3-8');self.w.results.setCurrentRow(2);self.w.verse_live()
        self.assertIn('John 2:5',self.w.live.current['reference'])
        self.w.previous_live();self.assertIn('John 2:4',self.w.live.current['reference'])
        self.w.query.setText('Psalm 23:1-2');QTest.keyClick(self.w.query,Qt.Key_Return)
        self.assertEqual(len(self.w.bible_flow_indices()),2)
        self.assertIn('23:1',self.w.live.current['reference'])
        self.w.next_live();self.assertIn('23:2',self.w.live.current['reference'])


    def test_passage_heading_and_each_number_on_live(self):
        from projection import slide_parts
        self.search('Gen 2:1-6')
        self.assertEqual(self.w.preview.current['passage_reference'],'Genesis 2:1-6')
        self.w.verse_live()
        for verse in range(1,7):
            item=self.w.live.current
            self.assertEqual(item['passage_reference'],'Genesis 2:1-6')
            self.assertTrue(slide_parts(item)[0].startswith(str(verse)+' '))
            self.assertEqual(item['body'],self.w.matches[verse-1]['text'])
            if verse<6:QTest.keyClick(self.w.results,Qt.Key_PageDown)
        self.w.previous_live()
        self.assertTrue(slide_parts(self.w.live.current)[0].startswith('5 '))
        self.qt.processEvents();canvas=self.w.live.content_layout.itemAt(0).widget();canvas.grab()
        self.assertGreater(canvas.passage_rect.height(),0)
        self.assertLessEqual(canvas.passage_rect.bottom(),canvas.body_rect.top())

    def test_number_survives_split_verse_and_queued_passage(self):
        from bible_core import paginate,passage_reference
        from projection import slide_parts
        row={'book':'Genesis','chapter':'2','verse':'1','text':' '.join(['word']*200)}
        pages=paginate([row],'KJV',1,150)
        self.assertGreater(len(pages),1)
        self.assertTrue(all(slide_parts(p)[0].startswith('1 ') for p in pages))
        self.assertEqual(' '.join(p['body'] for p in pages),row['text'])
        self.search('Genesis 2:1-6');self.w.add_passage()
        self.assertTrue(all(i['passage_reference']=='Genesis 2:1-6' for i in self.w.items))
        self.assertEqual(passage_reference(self.w.matches),'Genesis 2:1-6')

if __name__=='__main__':unittest.main()
